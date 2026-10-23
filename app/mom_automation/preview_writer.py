from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .excel_matcher import MatchResult
from .mail_parser import ParsedMailEvent


HEADERS = [
    "申请号/登记号",
    "规范化编号",
    "名称",
    "识别事件",
    "建议状态",
    "是否更新法律状态",
    "置信度",
    "邮件收到时间",
    "原邮件发送时间",
    "附件名",
    "匹配工作表",
    "匹配行",
    "匹配字段",
    "原状态",
    "建议新状态",
    "是否需要人工确认",
    "重复邮件数",
    "备注",
    "邮件主题",
]


def write_preview(
    events: list[ParsedMailEvent],
    matches: dict[str, MatchResult],
    output_path: Path,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "更新预览"
    ws.append(HEADERS)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    confirm_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for event in events:
        match = matches.get(event.dedupe_key, MatchResult())
        needs_confirmation = event.needs_confirmation or match.needs_confirmation
        note = "; ".join(n for n in [event.notes, match.note] if n)
        suggested_new_status = event.suggested_status if event.legal_status_update else match.current_status
        ws.append(
            [
                event.application_no,
                event.normalized_no,
                event.title or match.current_name,
                event.event,
                event.suggested_status,
                "是" if event.legal_status_update else "否",
                event.confidence,
                event.received_at,
                event.forwarded_at,
                "\n".join(event.attachment_names),
                match.sheet,
                match.row or "",
                match.matched_by,
                match.current_status,
                suggested_new_status,
                "是" if needs_confirmation else "否",
                event.duplicate_count,
                note,
                event.source_subject,
            ]
        )
        if needs_confirmation:
            for cell in ws[ws.max_row]:
                cell.fill = confirm_fill

    for col_idx, header in enumerate(HEADERS, 1):
        width = min(max(len(header) + 2, 12), 42)
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            value = row[0].value
            if value:
                width = min(max(width, len(str(value)) + 2), 50)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def write_log(events: list[ParsedMailEvent], matches: dict[str, MatchResult], output_path: Path) -> None:
    lines = [
        "妈妈办公自动化运行日志",
        "",
        f"解析事件数：{len(events)}",
        "",
    ]
    for event in events:
        match = matches.get(event.dedupe_key, MatchResult())
        lines.append(
            f"- {event.application_no or '未识别编号'} | {event.event or '未识别事件'} | "
            f"{event.suggested_status or '无建议状态'} | {match.note}"
        )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")

