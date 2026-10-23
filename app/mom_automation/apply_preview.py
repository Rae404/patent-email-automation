from __future__ import annotations

import shutil
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from .excel_matcher import STATUS_HEADER_KEYWORDS, find_header_row, select_columns


@dataclass
class ApplySummary:
    copied_to: Path
    updated_rows: int
    skipped_rows: int


def _read_preview_rows(preview_path: Path) -> list[dict[str, object]]:
    wb = load_workbook(preview_path, data_only=True)
    ws = wb.active
    headers = [str(c.value or "") for c in ws[1]]
    rows: list[dict[str, object]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rows.append({headers[i]: row[i] for i in range(len(headers))})
    return rows


def apply_preview(preview_path: Path, workbook_path: Path, output_path: Path) -> ApplySummary:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(workbook_path, output_path)
    preview_rows = _read_preview_rows(preview_path)

    wb = load_workbook(output_path)
    updated = 0
    skipped = 0

    for row in preview_rows:
        if row.get("是否需要人工确认") == "是" or row.get("是否更新法律状态") != "是":
            skipped += 1
            continue
        sheet_name = row.get("匹配工作表")
        row_idx = row.get("匹配行")
        new_status = row.get("建议新状态")
        if not sheet_name or not row_idx or not new_status or sheet_name not in wb.sheetnames:
            skipped += 1
            continue
        ws = wb[str(sheet_name)]
        _, headers = find_header_row(ws)
        status_cols = select_columns(headers, STATUS_HEADER_KEYWORDS)
        if not status_cols:
            skipped += 1
            continue
        ws.cell(int(row_idx), status_cols[0][1]).value = new_status
        updated += 1

    _write_update_record(wb, preview_rows)
    wb.save(output_path)
    return ApplySummary(output_path, updated, skipped)


def _write_update_record(wb, preview_rows: list[dict[str, object]]) -> None:
    sheet_name = "自动化更新记录"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    headers = [
        "申请号/登记号",
        "名称",
        "识别事件",
        "建议状态",
        "匹配工作表",
        "匹配行",
        "原状态",
        "建议新状态",
        "是否需要人工确认",
        "备注",
        "邮件主题",
    ]
    ws.append(headers)
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
    for row in preview_rows:
        ws.append([row.get(h, "") for h in headers])
    ws.freeze_panes = "A2"

