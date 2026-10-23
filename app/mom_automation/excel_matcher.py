from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .mail_parser import ParsedMailEvent
from .normalization import compact_text, normalize_identifier


MATCH_HEADER_KEYWORDS = [
    "专利号",
    "申请号",
    "软著登记号",
    "登记号",
    "公开(公告)号",
    "公开公告号",
]

STATUS_HEADER_KEYWORDS = ["法律状态/事件", "法律状态", "状态", "是否已取得授权证书"]
NAME_HEADER_KEYWORDS = ["知识产权名称", "专利名称", "软著名称", "名称"]


@dataclass
class MatchResult:
    sheet: str = ""
    row: int | None = None
    matched_by: str = ""
    current_status: str = ""
    current_name: str = ""
    needs_confirmation: bool = True
    note: str = "未提供主表，无法匹配"


def find_header_row(ws, max_scan_rows: int = 10) -> tuple[int, dict[str, int]]:
    for row_idx in range(1, min(ws.max_row, max_scan_rows) + 1):
        headers: dict[str, int] = {}
        for col_idx in range(1, ws.max_column + 1):
            value = ws.cell(row_idx, col_idx).value
            if value is not None:
                headers[compact_text(value)] = col_idx
        if any(any(k in h for k in MATCH_HEADER_KEYWORDS) for h in headers):
            return row_idx, headers
    return 1, {compact_text(ws.cell(1, col).value): col for col in range(1, ws.max_column + 1)}


def select_columns(headers: dict[str, int], keywords: Iterable[str]) -> list[tuple[str, int]]:
    result: list[tuple[str, int]] = []
    for header, col in headers.items():
        if any(keyword in header for keyword in keywords):
            result.append((header, col))
    return result


def match_events_to_workbook(events: list[ParsedMailEvent], workbook_path: Path | None) -> dict[str, MatchResult]:
    if workbook_path is None:
        return {event.dedupe_key: MatchResult() for event in events}

    wb = load_workbook(workbook_path, data_only=False)
    indexes = []
    for ws in wb.worksheets:
        header_row, headers = find_header_row(ws)
        id_cols = select_columns(headers, MATCH_HEADER_KEYWORDS)
        status_cols = select_columns(headers, STATUS_HEADER_KEYWORDS)
        name_cols = select_columns(headers, NAME_HEADER_KEYWORDS)
        if not id_cols:
            continue
        indexes.append((ws, header_row, id_cols, status_cols, name_cols))

    results: dict[str, MatchResult] = {}
    for event in events:
        candidates: list[MatchResult] = []
        for ws, header_row, id_cols, status_cols, name_cols in indexes:
            for row_idx in range(header_row + 1, ws.max_row + 1):
                for header, col_idx in id_cols:
                    cell_value = ws.cell(row_idx, col_idx).value
                    if not cell_value:
                        continue
                    if normalize_identifier(cell_value) == event.normalized_no:
                        current_status = _first_value(ws, row_idx, status_cols)
                        current_name = _first_value(ws, row_idx, name_cols)
                        candidates.append(
                            MatchResult(
                                sheet=ws.title,
                                row=row_idx,
                                matched_by=header,
                                current_status=current_status,
                                current_name=current_name,
                                needs_confirmation=False,
                                note="匹配成功",
                            )
                        )
        if len(candidates) == 1:
            result = candidates[0]
            result.needs_confirmation = event.needs_confirmation
            if event.needs_confirmation:
                result.note = "匹配成功，但事件类型建议人工确认"
            results[event.dedupe_key] = result
        elif len(candidates) > 1:
            results[event.dedupe_key] = MatchResult(
                needs_confirmation=True,
                note=f"同一申请号匹配到多行：{len(candidates)}",
            )
        else:
            results[event.dedupe_key] = MatchResult(
                needs_confirmation=True,
                note="主表中未匹配到申请号/登记号",
            )
    return results


def _first_value(ws, row_idx: int, cols: list[tuple[str, int]]) -> str:
    for _, col_idx in cols:
        value = ws.cell(row_idx, col_idx).value
        if value not in (None, ""):
            return str(value)
    return ""

