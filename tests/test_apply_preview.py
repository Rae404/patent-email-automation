import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "app"))

from mom_automation.apply_preview import apply_preview
from mom_automation.preview_writer import write_preview
from mom_automation.excel_matcher import match_events_to_workbook
from mom_automation.mail_parser import load_samples_json, parse_samples


class ApplyPreviewTests(unittest.TestCase):
    def test_apply_preview_updates_only_safe_rows_and_adds_record_sheet(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            workbook = tmp_path / "main.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "总表"
            ws.append(["专利号", "知识产权名称", "法律状态/事件"])
            ws.append(["209910000001.1", "虚构示例方法", "受理"])
            wb.save(workbook)

            samples = load_samples_json(Path(__file__).resolve().parents[1] / "samples" / "emails.json")
            events = parse_samples(samples)
            matches = match_events_to_workbook(events, workbook)
            preview = tmp_path / "preview.xlsx"
            write_preview(events, matches, preview)

            out = tmp_path / "updated.xlsx"
            summary = apply_preview(preview, workbook, out)
            self.assertGreaterEqual(summary.updated_rows, 1)

            updated = load_workbook(out)
            self.assertEqual(updated["总表"].cell(2, 3).value, "授权")
            self.assertIn("自动化更新记录", updated.sheetnames)


if __name__ == "__main__":
    unittest.main()
